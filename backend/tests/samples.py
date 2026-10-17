"""Builders for real synthetic test documents (PDF, DOCX, TXT) and deliberately broken files.

Run as a script to write sample files for manual testing:
    .venv\\Scripts\\python -m tests.samples ..\\samples
"""
from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

INVOICE_TEXT = """Northwind Traders Ltd
INVOICE
Invoice Number: INV-48213
Invoice Date: 4 March 2025
Bill To: Harbor Dental Clinic, 22 Bay Street, Portland
Description: Annual software licence for practice management, 5 seats.
Subtotal $4,200.00. Sales tax 8% $336.00. Total amount due $4,536.00.
Payment is due within 14 days of the invoice date. Please quote the invoice number with your payment.
Late payments may incur an interest charge of 1.5% per month.
Questions about this bill: accounts@northwind.example or +1 503 555 0199."""

CONTRACT_TEXT = """CONSULTING AGREEMENT
This Agreement is entered into on 1 February 2025 by and between Bluepeak Software Inc. ("Client") and Maria Lopez ("Consultant").
The Consultant shall provide data engineering services as described in Schedule A.
The term of this Agreement shall commence on the effective date and continue for twelve months.
In consideration of the services, the Client shall pay a fee of $9,000 per month, invoiced monthly.
Either party may terminate this Agreement with thirty days written notice.
Confidential Information shall not be disclosed to any third party without prior written consent, and these obligations survive termination.
This Agreement shall be governed by the laws of the State of Texas.
IN WITNESS WHEREOF, the parties have executed this Agreement as of the date first written above."""

RESUME_TEXT = """Aisha Khan
Senior Data Analyst
Professional summary: data analyst with seven years of experience turning complex data into clear business insight.
Work experience: Senior Data Analyst at Summit Capital Group, Toronto, 2021 to present. Built forecasting dashboards used by 40 managers.
Data Analyst at Riverside Media, 2017 to 2021. Automated weekly reporting with Python and SQL, saving ten hours per week.
Education: Bachelor of Science in Statistics, University of Toronto, 2017.
Skills: Python, SQL, Tableau, stakeholder management, statistical modelling.
Certifications: Google Data Analytics Professional Certificate.
References available upon request."""

MEDICAL_TEXT = """Patient Consultation Note
Patient: James Carter. Date of visit: 12 March 2025.
Chief complaint: persistent fatigue and shortness of breath on exertion for three weeks.
Examination: blood pressure 142/90 mmHg, heart rate 88 bpm. Chest clear on auscultation.
Laboratory findings: haemoglobin 10.1 g/dL, below the reference range; ferritin low.
Diagnosis: iron deficiency anaemia.
Plan: start ferrous sulfate 200 mg twice daily, repeat complete blood count in six weeks, dietary advice given.
Follow-up appointment recommended in 6 weeks. Attending physician: Dr. Chen Wei, Internal Medicine."""

LEGAL_TEXT = """FINAL NOTICE BEFORE LEGAL ACTION
Take notice that our client, Greenfield Logistics, has instructed us to issue this notice.
Despite repeated reminders, you have failed to pay the outstanding sum of $12,300 under the agreement dated 2 June 2024.
You are hereby required to pay the full amount within 7 days of the date of this notice.
Failing compliance, our client will initiate legal proceedings before the District Court without further notice.
You will be liable for all costs and consequences arising from legal action. Please treat this matter as urgent."""

OUT_OF_SCOPE_TEXT = """Grandma's apple crumble.
Peel and slice six apples and place them in a baking dish with cinnamon and brown sugar.
Rub butter into flour and oats until it resembles breadcrumbs, then sprinkle over the fruit.
Bake for forty minutes until golden and bubbling at the edges.
Serve warm with custard or vanilla ice cream. Leftovers keep in the fridge for two days."""

SAMPLES = {
    "invoice": INVOICE_TEXT,
    "contract": CONTRACT_TEXT,
    "resume": RESUME_TEXT,
    "medical_report": MEDICAL_TEXT,
    "legal_notice": LEGAL_TEXT,
}


def make_txt(text: str) -> bytes:
    return text.encode("utf-8")


def make_pdf(text: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in text.split("\n"):
        safe = line.encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(0, 6, safe, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def make_image_only_pdf() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_fill_color(30, 40, 90)
    pdf.rect(20, 20, 120, 60, style="F")
    return bytes(pdf.output())


def make_encrypted_pdf(text: str) -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(make_pdf(text))).pages:
        writer.add_page(page)
    writer.encrypt(user_password="secret-pass", owner_password="owner-pass")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx(text: str, with_table: bool = True) -> bytes:
    import docx

    document = docx.Document()
    lines = text.split("\n")
    document.add_heading(lines[0], level=1)
    for line in lines[1:]:
        document.add_paragraph(line)
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Prepared by"
        table.cell(0, 1).text = "Document control"
        table.cell(1, 0).text = "Status"
        table.cell(1, 1).text = "Issued"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_corrupt_docx() -> bytes:
    """A valid zip archive that is not a Word document."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "not a word document")
    return buf.getvalue()


def make_docx_bomb() -> bytes:
    """A zip that declares a huge, highly compressible member inside a DOCX-shaped archive."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", "A" * (70 * 1024 * 1024))
    return buf.getvalue()


def write_samples(out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    written = []
    builders = {"pdf": make_pdf, "docx": make_docx, "txt": make_txt}
    plan = [("invoice", "pdf"), ("contract", "docx"), ("resume", "pdf"), ("medical_report", "txt"),
            ("legal_notice", "docx"), ("invoice", "txt")]
    for name, ext in plan:
        path = out_dir / f"sample_{name}.{ext}"
        path.write_bytes(builders[ext](SAMPLES[name]))
        written.append(path)
    extra = {
        "sample_out_of_scope_recipe.txt": make_txt(OUT_OF_SCOPE_TEXT),
        "invalid_scanned_image_only.pdf": make_image_only_pdf(),
        "invalid_not_really_a.pdf": b"This is plain text pretending to be a PDF.",
    }
    for name, data in extra.items():
        (out_dir / name).write_bytes(data)
        written.append(out_dir / name)
    return written


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "samples")
    for p in write_samples(target):
        print(p)
